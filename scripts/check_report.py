#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""check_report.py — Verificacion final del informe antes de entregarlo.

Hace de una sola pasada las dos comprobaciones del Paso 8 del skill:

  1. Que el .docx abra sin corrupcion (si esto falla, Word lo va a rechazar).
  2. Cuantos recuadros "[ PEGAR AQUI ]" quedaron y en que seccion esta cada uno.

El conteo recorre parrafos Y celdas de tabla. Esto importa: los recuadros del
informe viven dentro de tablas, y `Document.paragraphs` no entra en las tablas,
asi que contar solo sobre parrafos devuelve 0 aunque haya recuadros pendientes.

Opcionalmente valida tambien el metadata.json y lista los campos que quedaron
como "(completar ...)", que es lo otro que hay que reportarle al usuario.

Uso:
  python scripts/check_report.py work/informe.docx [--metadata work/metadata.json]

Sale con codigo 1 si el .docx no abre. Los recuadros pendientes NO son un error:
son informacion para el mensaje de entrega.
"""
from __future__ import print_function

import argparse
import io
import json
import os
import sys

# Una consola de Windows sin UTF-8 (cmd.exe, PowerShell con la codepage por
# defecto) no puede imprimir cualquier caracter Unicode; si alguno se cuela en
# un print() futuro, la corrida termina en UnicodeEncodeError en vez de
# terminar la verificacion. reconfigure() con errors='replace' cambia eso por
# un simple '?' en pantalla, nunca un crash. No falla en Python viejo: si
# reconfigure no existe, sigue con el comportamiento por defecto.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(errors='replace')
    except AttributeError:
        pass

MARKER = 'PEGAR'


def ruta_larga(p):
    """Windows corta las rutas en 260 caracteres y python-docx falla al abrirlas.

    El prefijo \\\\?\\ desactiva ese limite. Solo se aplica en Windows y cuando la
    ruta ya es larga, para no ensuciar la salida en el caso normal.
    """
    ap = os.path.abspath(p)
    if os.name == 'nt' and len(ap) > 250 and not ap.startswith('\\\\?\\'):
        return '\\\\?\\' + ap
    return ap


def iter_cells(doc):
    """Devuelve (texto, contexto) de cada celda de tabla del documento."""
    for ti, table in enumerate(doc.tables):
        for row in table.rows:
            for cell in row.cells:
                yield cell.text, 'tabla %d' % (ti + 1)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('docx', help='ruta del informe .docx a verificar')
    ap.add_argument('--metadata', help='metadata.json de la corrida, para listar los (completar ...)')
    args = ap.parse_args()

    try:
        from docx import Document
    except ImportError:
        print('ERROR: falta python-docx. Ver README.md (pip install python-docx).', file=sys.stderr)
        return 1

    try:
        doc = Document(ruta_larga(args.docx))
    except Exception as e:
        print('ERROR: el .docx no abre, esta corrupto y Word lo va a rechazar.', file=sys.stderr)
        print('  %s: %s' % (type(e).__name__, e), file=sys.stderr)
        return 1

    print('OK: %s abre correctamente (%d parrafos, %d tablas)'
          % (args.docx, len(doc.paragraphs), len(doc.tables)))

    def es_titulo(p):
        # p.style puede venir en None, y el nombre del estilo cambia con el
        # idioma de Word ("Heading 1" / "Titulo 1"), asi que se toleran ambos.
        style = getattr(p, 'style', None)
        nombre = (getattr(style, 'name', None) or '').lower()
        return nombre.startswith('heading') or nombre.startswith('tit')

    pendientes = []
    ultimo_titulo = '(inicio del documento)'
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt and es_titulo(p):
            ultimo_titulo = txt
        if MARKER in txt:
            pendientes.append((ultimo_titulo, txt))
    for txt, ctx in iter_cells(doc):
        if MARKER in txt:
            pendientes.append((ctx, ' '.join(txt.split())))

    print('')
    if pendientes:
        print('Recuadros pendientes de pegar a mano: %d' % len(pendientes))
        for donde, txt in pendientes:
            corto = txt if len(txt) <= 90 else txt[:87] + '...'
            print('  - [%s] %s' % (donde, corto))
    else:
        print('Recuadros pendientes de pegar a mano: 0 (informe completo)')

    if args.metadata:
        try:
            meta = json.load(io.open(ruta_larga(args.metadata), encoding='utf-8'))
        except Exception as e:
            print('')
            print('AVISO: no se pudo leer %s (%s)' % (args.metadata, e))
        else:
            faltantes = []

            def scan(prefix, value):
                if isinstance(value, dict):
                    for k, v in value.items():
                        scan('%s.%s' % (prefix, k) if prefix else k, v)
                elif isinstance(value, list):
                    for i, v in enumerate(value):
                        scan('%s[%d]' % (prefix, i), v)
                elif isinstance(value, str) and 'completar' in value.lower():
                    faltantes.append((prefix, value))

            scan('', meta)

            # tipo_proyecto es un campo binario: "migracion" o
            # "implementacion_nueva". Un valor vacio ("") no contiene la
            # palabra "completar", asi que el scan() generico de arriba
            # nunca lo detecta -- y build_report.js, si no se valida, elige
            # un tipo en silencio (ver el fix en ese archivo). Se chequea
            # aparte porque es el unico campo donde "vacio" no es lo mismo
            # que "inofensivo": cambia el titulo del documento entero.
            tp = meta.get('tipo_proyecto')
            if tp not in ('migracion', 'implementacion_nueva'):
                faltantes.append(('tipo_proyecto', repr(tp) + ' (debe ser "migracion" o "implementacion_nueva")'))

            print('')
            if faltantes:
                print('Campos de metadata sin completar: %d' % len(faltantes))
                for campo, valor in faltantes:
                    print('  - %s = %s' % (campo, valor))
            else:
                print('Campos de metadata sin completar: 0')

    return 0


if __name__ == '__main__':
    sys.exit(main())
